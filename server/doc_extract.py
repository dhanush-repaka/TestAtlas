"""Extracts plain text from uploaded documents in whatever format a real
project's documentation actually happens to be in -- not just plain
text/Markdown. All extraction is local, pure-Python, and free (no API calls),
consistent with the rest of this project's cost constraints.
"""
from __future__ import annotations

import csv
import io

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".docx", ".pptx", ".xlsx", ".pdf"}

_MAX_CHARS = 300_000  # generous, but keeps a pathological huge file from blowing up
                       # the DB row / gap-analysis context size


class UnsupportedDocumentError(ValueError):
    pass


def _extension_of(filename: str) -> str:
    return ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""


def extract_text(filename: str, data: bytes) -> str:
    """Returns the document's text content. Raises UnsupportedDocumentError
    (safe to show directly to the user) for an unrecognized extension, or if
    extraction produced nothing usable."""
    ext = _extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{ext or filename}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if ext in (".txt", ".md", ".markdown"):
        text = data.decode("utf-8", errors="replace")
    elif ext == ".csv":
        text = _extract_csv(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    elif ext == ".pptx":
        text = _extract_pptx(data)
    elif ext == ".xlsx":
        text = _extract_xlsx(data)
    else:  # .pdf
        text = _extract_pdf(data)

    text = text.strip()
    if not text:
        raise UnsupportedDocumentError("Couldn't find any text in that file -- it may be empty, image-only, or corrupted.")
    if len(text) > _MAX_CHARS:
        text = text[:_MAX_CHARS] + "\n\n[...truncated -- document exceeded the size this feature keeps in one document...]"
    return text


def _extract_docx(data: bytes) -> str:
    import docx  # lazy import: only needed when a .docx is actually uploaded

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pptx(data: bytes) -> str:
    from pptx import Presentation  # lazy import

    presentation = Presentation(io.BytesIO(data))
    slides_text = []
    for i, slide in enumerate(presentation.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        if parts:
            slides_text.append(f"--- Slide {i} ---\n" + "\n".join(parts))
    return "\n\n".join(slides_text)


def _normalize_cell(value: str) -> str:
    """Collapses any run of whitespace -- including an embedded newline, routine
    inside a real ADO "Acceptance Criteria" field exported as a multi-line
    bulleted list -- down to a single space, and strips the ends. Downstream
    (server/work_item_import.py) reads this text one LINE per work-item row,
    so a cell that keeps an internal newline would silently fragment its own
    row into several, truncating that row's own content and producing
    orphaned garbage lines that don't match any real row shape."""
    return " ".join(value.split())


def _extract_csv(data: bytes) -> str:
    """Same "cell | cell | cell" per-row shape _extract_xlsx produces, so
    anything downstream that reads a pipe-delimited table (server/
    work_item_import.py's header/row detection, for an ADO work-item export)
    doesn't care whether the original file was a real spreadsheet or a plain
    CSV -- one format either way. Uses csv.reader, not a naive split(','), so
    a quoted field containing a comma is handled correctly, and
    _normalize_cell handles one containing a newline. Keeps every cell,
    including a blank one, rather than dropping it -- work_item_import.py
    looks a row's fields up by the HEADER's column index, so a blank
    "Description" (routine on a real ADO export -- plenty of stories skip
    it) must not shift every column after it out of position, silently
    misattributing a later row's Acceptance Criteria to the wrong field. A
    row with nothing in it at all is still skipped, same as before."""
    text = data.decode("utf-8-sig", errors="replace")  # utf-8-sig: strips an Excel-added BOM if present
    rows_text = []
    for row in csv.reader(io.StringIO(text)):
        cells = [_normalize_cell(c) for c in row]
        if any(cells):
            rows_text.append(" | ".join(cells))
    return "\n".join(rows_text)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl  # lazy import

    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    sheets_text = []
    for ws in workbook.worksheets:
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            # Same positional-alignment and embedded-newline fix as _extract_csv:
            # a blank cell stays a blank cell in the row (never dropped, which
            # would shift later columns out of position), and a cell's own
            # internal newline is collapsed so it can never fragment its row.
            cells = [_normalize_cell(str(c)) if c is not None else "" for c in row]
            if any(cells):
                rows_text.append(" | ".join(cells))
        if rows_text:
            sheets_text.append(f"--- Sheet: {ws.title} ---\n" + "\n".join(rows_text))
    return "\n\n".join(sheets_text)


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader  # lazy import

    reader = PdfReader(io.BytesIO(data))
    pages_text = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(t for t in pages_text if t)
